#!/usr/bin/env python3

from __future__ import annotations

import argparse
import re
import warnings
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

from docx_utils import apply_font_size_to_runs, apply_highlight_to_runs
from generate_subs import is_full_line_comment
from style_tokens import BODY_TEXT_SIZE_PT
from template_styles import ensure_base_styles


DEFAULT_TEMPLATE_PATH = Path("templates/news_template.docx")
TIME_TOKEN = r"\d{1,2}:\d{2}:\d{2}:\d{2}"
TIMED_LINE_RE = re.compile(
    rf"^(?P<start>{TIME_TOKEN})(?P<between>[\t \u3000]+)"
    rf"(?P<end>{TIME_TOKEN})(?:(?P<after>[\t \u3000]+)(?P<text>.*))?$"
)
SUPER_AS_SUBTITLE_RE = re.compile(
    r"super.*字幕方式.*copy.*字幕檔", re.IGNORECASE
)
NO_TRANSLATION_RE = re.compile(r"(?:以下)?不用翻譯|do not translate", re.IGNORECASE)
STAR_MARKER_RE = re.compile(r"\s*\*\s*$")
SHARING_SKY_DELIVERY_TITLE_RE = re.compile(
    r"(?P<title>風月同天第.+?\(\s*12分版\s*\))"
)


class DocumentKind(str, Enum):
    SUBTITLE = "subtitle"
    SUPER = "super"


@dataclass
class DocsBlock:
    timecode: str | None = None
    source_lines: list[str] = field(default_factory=list)
    english_lines: list[str] = field(default_factory=list)
    cyan: bool = False
    no_translation: bool = False
    inline_translation: bool = False
    highlight_marker: bool = False
    yellow: bool = False
    timestamp_inline_translation: bool = False


@dataclass
class ParsedDocs:
    kind: DocumentKind
    blocks: list[DocsBlock]
    blank_after: set[int] = field(default_factory=set)


def _decode_input(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16")
    for encoding in ("utf-8-sig", "big5", "cp950", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UnicodeError(f"Unable to decode input file: {path}")


def _split_bilingual(text: str) -> tuple[str | None, str | None]:
    if "//" not in text:
        return text, None
    source, english = text.split("//", 1)
    return (source if source else None), (english if english else None)


def _timecode_and_text(line: str) -> tuple[str, str | None] | None:
    match = TIMED_LINE_RE.fullmatch(line)
    if match is None:
        return None
    timecode = f"{match.group('start')}\t{match.group('end')}"
    return timecode, match.group("text")


def _timed_line_and_marker(line: str) -> tuple[tuple[str, str | None] | None, bool]:
    has_marker = STAR_MARKER_RE.search(line) is not None
    candidate = STAR_MARKER_RE.sub("", line) if has_marker else line
    timed = _timecode_and_text(candidate)
    if timed is None:
        return _timecode_and_text(line), False
    return timed, has_marker


def _apply_star_highlight_range(blocks: list[DocsBlock]) -> None:
    marker_indices = [
        index for index, block in enumerate(blocks) if block.highlight_marker
    ]
    for pair_start in range(0, len(marker_indices) - 1, 2):
        start, end = marker_indices[pair_start : pair_start + 2]
        for block in blocks[start : end + 1]:
            block.yellow = True


def _parse_subtitle(lines: list[str]) -> ParsedDocs:
    blocks: list[DocsBlock] = []
    blank_after: set[int] = set()
    current: DocsBlock | None = None

    for line in lines:
        timed, highlight_marker = _timed_line_and_marker(line)
        if timed is not None:
            current = DocsBlock(
                timecode=timed[0],
                inline_translation=timed[1] is not None and "//" in timed[1],
                highlight_marker=highlight_marker,
            )
            blocks.append(current)
            if timed[1] is not None:
                source, english = _split_bilingual(timed[1])
                if source is not None:
                    current.source_lines.append(source)
                if english is not None:
                    current.english_lines.append(english)
            continue
        if not line:
            if blocks:
                blank_after.add(len(blocks) - 1)
            current = None
            continue
        if current is None:
            current = DocsBlock()
            blocks.append(current)
        source, english = _split_bilingual(line)
        if source is not None:
            current.source_lines.append(source)
        if english is not None:
            current.english_lines.append(english)

    _apply_star_highlight_range(blocks)
    return ParsedDocs(DocumentKind.SUBTITLE, blocks, blank_after)


def _parse_super(lines: list[str]) -> ParsedDocs:
    blocks: list[DocsBlock] = []
    blank_after: set[int] = set()
    current: DocsBlock | None = None
    no_translation = False
    cyan_next = False

    def finish_block() -> None:
        nonlocal current
        if current is None:
            return
        current = None

    for line in lines:
        if not line:
            finish_block()
            if blocks:
                blank_after.add(len(blocks) - 1)
            continue

        timed, highlight_marker = _timed_line_and_marker(line)
        if timed is not None:
            finish_block()
            current = DocsBlock(
                timecode=timed[0],
                no_translation=no_translation,
                cyan=cyan_next,
                inline_translation=timed[1] is not None and "//" in timed[1],
                highlight_marker=highlight_marker,
                timestamp_inline_translation=(
                    timed[1] is not None and "//" in timed[1]
                ),
            )
            cyan_next = False
            blocks.append(current)
            content = timed[1]
            if content is None:
                continue
        else:
            content = line
            if current is None:
                current = DocsBlock(
                    no_translation=no_translation,
                    cyan=cyan_next,
                )
                cyan_next = False
                blocks.append(current)

        if "//" in content:
            current.inline_translation = True
        source, english = _split_bilingual(content)
        if source is not None:
            current.source_lines.append(source)
        if english is not None:
            current.english_lines.append(english)

        if source and SUPER_AS_SUBTITLE_RE.search(source):
            current.cyan = True
            cyan_next = True
        if source and NO_TRANSLATION_RE.search(source):
            current.no_translation = True
            no_translation = True

    finish_block()

    _apply_star_highlight_range(blocks)
    return ParsedDocs(DocumentKind.SUPER, blocks, blank_after)


def parse_docs_text(text: str, kind: DocumentKind) -> ParsedDocs:
    lines = [
        line
        for line in text.lstrip("\ufeff").splitlines()
        if not is_full_line_comment(line)
    ]
    if kind is DocumentKind.SUBTITLE:
        return _parse_subtitle(lines)
    return _parse_super(lines)


def detect_document_kind(path: Path, text: str) -> DocumentKind:
    name = path.name.casefold()
    if "super" in name:
        return DocumentKind.SUPER
    if "chus字幕" in name or "英文字幕" in name or "subtitle" in name:
        return DocumentKind.SUBTITLE

    timed_lines = [
        timed
        for line in text.splitlines()
        if (timed := _timecode_and_text(line)) is not None
    ]
    with_inline_text = sum(timed[1] is not None for timed in timed_lines)
    if timed_lines and with_inline_text == len(timed_lines):
        return DocumentKind.SUBTITLE
    if timed_lines and with_inline_text == 0:
        return DocumentKind.SUPER
    raise ValueError(
        f"Unable to determine Docs input type for '{path.name}'; "
        "use --kind subtitle or --kind super."
    )


def _remove_initial_paragraph(doc: Document) -> None:
    for paragraph in list(doc.paragraphs):
        paragraph._element.getparent().remove(paragraph._element)


def _add_paragraph(
    doc: Document,
    text: str,
    kind: DocumentKind,
    *,
    cyan: bool = False,
    yellow: bool = False,
):
    text = re.sub(r"\s+#\s*$", "", text)
    paragraph = doc.add_paragraph()
    if text or cyan:
        paragraph.add_run(text)
    apply_font_size_to_runs(paragraph, font_size_pt=BODY_TEXT_SIZE_PT)
    if cyan:
        apply_highlight_to_runs(
            paragraph,
            highlight_color=WD_COLOR_INDEX.TURQUOISE,
        )
    elif yellow:
        apply_highlight_to_runs(
            paragraph,
            highlight_color=WD_COLOR_INDEX.YELLOW,
        )
    return paragraph


def _validate(parsed: ParsedDocs) -> None:
    if parsed.kind is DocumentKind.SUBTITLE:
        too_long = [
            line
            for block in parsed.blocks
            for line in block.english_lines
            if len(line) > 54
        ]
        if too_long:
            warnings.warn(
                f"{len(too_long)} English subtitle line(s) exceed 54 characters; "
                "content was preserved unchanged.",
                UserWarning,
                stacklevel=2,
            )
        return

    invalid = [
        line
        for block in parsed.blocks
        if not block.timestamp_inline_translation
        for line in block.english_lines
        if "." in line or "(" in line or ")" in line
    ]
    if invalid:
        warnings.warn(
            f"{len(invalid)} English SUPER line(s) contain periods or parentheses; "
            "content was preserved unchanged.",
            UserWarning,
            stacklevel=2,
        )


def _resolve_template_path(template_path: Path) -> Path:
    if template_path.is_absolute() or template_path.exists():
        return template_path
    return Path(__file__).resolve().parent / template_path


def _render(parsed: ParsedDocs, output_path: Path, template_path: Path) -> Path:
    doc = Document(str(_resolve_template_path(template_path)))
    _remove_initial_paragraph(doc)
    ensure_base_styles(doc)

    for index, block in enumerate(parsed.blocks):
        yellow = block.yellow and not block.inline_translation and not block.no_translation
        if block.timecode is not None:
            if parsed.kind is DocumentKind.SUBTITLE and block.source_lines:
                _add_paragraph(
                    doc,
                    f"{block.timecode}\t{block.source_lines[0]}",
                    parsed.kind,
                    cyan=block.cyan,
                    yellow=yellow,
                )
                remaining_source = block.source_lines[1:]
            else:
                _add_paragraph(
                    doc,
                    block.timecode,
                    parsed.kind,
                    cyan=block.cyan,
                    yellow=yellow,
                )
                remaining_source = block.source_lines
        else:
            remaining_source = block.source_lines

        for line in remaining_source:
            _add_paragraph(
                doc, line, parsed.kind, cyan=block.cyan, yellow=yellow
            )
        for line in block.english_lines:
            _add_paragraph(
                doc, line, parsed.kind, cyan=block.cyan, yellow=yellow
            )

        if parsed.kind is DocumentKind.SUPER and index in parsed.blank_after:
            _add_paragraph(
                doc, "", parsed.kind, cyan=block.cyan, yellow=yellow
            )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def default_output_path(input_path: Path, kind: DocumentKind) -> Path:
    delivery_title = SHARING_SKY_DELIVERY_TITLE_RE.search(input_path.stem)
    if delivery_title is not None:
        suffix = "字幕_final" if kind is DocumentKind.SUBTITLE else "super_final"
        return Path("output") / f"{delivery_title.group('title')}_{suffix}.docx"
    return Path("output") / input_path.with_suffix(".docx").name


def discover_input_paths(directory: Path) -> list[Path]:
    inputs: list[Path] = []
    for path in sorted(directory.glob("*.txt")):
        name = path.name.casefold()
        if name.startswith("~") or name.endswith(".baseline.txt"):
            continue
        try:
            detect_document_kind(path, "")
        except ValueError:
            continue
        inputs.append(path)
    return inputs


def generate_docs(
    input_path: Path,
    *,
    output_path: Path | None = None,
    kind: DocumentKind | None = None,
    template_path: Path = DEFAULT_TEMPLATE_PATH,
) -> Path:
    text = _decode_input(input_path)
    resolved_kind = kind or detect_document_kind(input_path, text)
    parsed = parse_docs_text(text, resolved_kind)
    _validate(parsed)
    return _render(
        parsed,
        output_path or default_output_path(input_path, resolved_kind),
        template_path,
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        prog="gen-docs",
        description="Generate a 風月同天 subtitle or SUPER working DOCX from TXT."
    )
    parser.add_argument(
        "--input",
        default="",
        help=(
            "Subtitle or SUPER TXT file. When omitted, generate all recognized "
            "non-baseline TXT files in the current directory."
        ),
    )
    parser.add_argument(
        "--output",
        default="",
        help=(
            "Output DOCX path. Default: ./output/ using the editor delivery name "
            "for recognized 風月同天 files, otherwise the input basename."
        ),
    )
    parser.add_argument(
        "--template",
        default=str(DEFAULT_TEMPLATE_PATH),
        help="Synchronized project DOCX template.",
    )
    parser.add_argument(
        "--kind",
        choices=("auto", *[kind.value for kind in DocumentKind]),
        default="auto",
        help="Input type; defaults to reliable filename/content detection.",
    )
    args = parser.parse_args()

    kind = None if args.kind == "auto" else DocumentKind(args.kind)
    try:
        input_paths = (
            [Path(args.input)]
            if args.input
            else discover_input_paths(Path.cwd())
        )
        if not input_paths:
            raise ValueError(
                "No recognized non-baseline subtitle or SUPER TXT files "
                "were found in the current directory."
            )
        if args.output and len(input_paths) > 1:
            raise ValueError("--output can only be used with one --input file.")
        for input_path in input_paths:
            output = generate_docs(
                input_path,
                output_path=Path(args.output) if args.output else None,
                kind=kind,
                template_path=Path(args.template),
            )
            print(output)
    except (FileNotFoundError, UnicodeError, ValueError) as exc:
        raise SystemExit(f"[error] {exc}") from exc


if __name__ == "__main__":
    main()
