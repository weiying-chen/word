from pathlib import Path
import zipfile

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

import generate_docs as generate_docs_module
from generate_docs import (
    DocumentKind,
    default_output_path,
    detect_document_kind,
    generate_docs,
    parse_docs_text,
)
from style_tokens import (
    BODY_TEXT_SIZE_PT,
    DEFAULT_DOCX_ASCII_FONT_NAME,
    DEFAULT_DOCX_EAST_ASIA_FONT_NAME,
)


def test_gen_docs_wrapper_uses_project_virtual_environment() -> None:
    wrapper = Path(__file__).resolve().parents[1] / "gen-docs"

    assert wrapper.stat().st_mode & 0o111
    assert ".venv/bin/python" in wrapper.read_text(encoding="utf-8")


def test_subtitle_parser_preserves_timecodes_and_splits_bilingual_text() -> None:
    source = (
        "00:01:02:03\t00:01:04:05\t中文內容//Existing English\n"
        "00:01:04:05\t00:01:06:07\t尚待翻譯\n"
    )

    parsed = parse_docs_text(source, DocumentKind.SUBTITLE)

    assert parsed.blocks[0].timecode == "00:01:02:03\t00:01:04:05"
    assert parsed.blocks[0].source_lines == ["中文內容"]
    assert parsed.blocks[0].english_lines == ["Existing English"]
    assert parsed.blocks[1].source_lines == ["尚待翻譯"]
    assert parsed.blocks[1].english_lines == []


def test_super_parser_normalizes_full_width_timecode_separator_to_tab() -> None:
    source = (
        "00:01:39:08　00:01:43:02\n"
        "2022年八月 美國\n"
        "俄亥俄州哥倫布市\n\n"
        "來源說明\n"
    )

    parsed = parse_docs_text(source, DocumentKind.SUPER)

    assert parsed.blocks[0].timecode == "00:01:39:08\t00:01:43:02"
    assert parsed.blocks[0].source_lines == ["2022年八月 美國", "俄亥俄州哥倫布市"]
    assert parsed.blocks[1].timecode is None
    assert parsed.blocks[1].source_lines == ["來源說明"]


def test_super_parser_keeps_consecutive_inline_english_timecodes() -> None:
    source = (
        "00:00:01:00\t00:00:02:00\t//First English line\n"
        "00:00:02:00\t00:00:03:00\t//Second English line\n"
    )

    parsed = parse_docs_text(source, DocumentKind.SUPER)

    assert [block.timecode for block in parsed.blocks] == [
        "00:00:01:00\t00:00:02:00",
        "00:00:02:00\t00:00:03:00",
    ]
    assert [block.english_lines for block in parsed.blocks] == [
        ["First English line"],
        ["Second English line"],
    ]


@pytest.mark.parametrize(
    ("filename", "expected"),
    [
        ("episode_chus字幕.txt", DocumentKind.SUBTITLE),
        ("episode_英文字幕.txt", DocumentKind.SUBTITLE),
        ("episode_super.txt", DocumentKind.SUPER),
    ],
)
def test_document_kind_detection_from_filename(
    filename: str, expected: DocumentKind
) -> None:
    assert detect_document_kind(Path(filename), "sample") is expected


def test_document_kind_detection_rejects_ambiguous_input() -> None:
    with pytest.raises(ValueError, match="Unable to determine"):
        detect_document_kind(Path("episode.txt"), "unstructured text")


@pytest.mark.parametrize("kind", [DocumentKind.SUBTITLE, DocumentKind.SUPER])
def test_default_output_path_uses_output_directory_and_preserves_basename(
    kind: DocumentKind,
) -> None:
    source = Path("incoming") / "episode_chus字幕.txt"

    assert default_output_path(source, kind) == Path("output/episode_chus字幕.docx")


@pytest.mark.parametrize(
    ("source_name", "kind", "expected_name"),
    [
        (
            "TO編譯-風月同天第4集_獅子山篇(12分版)_chus字幕(確定)(備註).txt",
            DocumentKind.SUBTITLE,
            "風月同天第4集_獅子山篇(12分版)_字幕_final.docx",
        ),
        (
            "TO編譯-風月同天第4集_獅子山篇(12分版)_super(確定)(備註).txt",
            DocumentKind.SUPER,
            "風月同天第4集_獅子山篇(12分版)_super_final.docx",
        ),
    ],
)
def test_default_output_path_uses_editor_delivery_name_for_sharing_sky(
    source_name: str,
    kind: DocumentKind,
    expected_name: str,
) -> None:
    assert default_output_path(Path(source_name), kind) == Path("output") / expected_name


def test_subtitle_generation_places_english_below_chinese_and_warns_on_length(
    tmp_path: Path,
) -> None:
    source = tmp_path / "episode_chus字幕.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\t中文//" + "x" * 55 + "\n",
        encoding="utf-8",
    )
    output = tmp_path / "subtitle.docx"

    with pytest.warns(UserWarning, match="54 characters"):
        generate_docs(source, output_path=output)

    doc = Document(output)
    assert [paragraph.text for paragraph in doc.paragraphs] == [
        "00:00:00:00\t00:00:02:00\t中文",
        "x" * 55,
    ]
    assert all(paragraph.paragraph_format.line_spacing is None for paragraph in doc.paragraphs)


def test_subtitle_generation_omits_blank_translation_line(tmp_path: Path) -> None:
    source = tmp_path / "episode_chus字幕.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\t尚待翻譯\n",
        encoding="utf-8",
    )
    output = tmp_path / "subtitle.docx"

    generate_docs(source, output_path=output)

    assert [paragraph.text for paragraph in Document(output).paragraphs] == [
        "00:00:00:00\t00:00:02:00\t尚待翻譯",
    ]


def test_generation_removes_trailing_workflow_hash_markers(tmp_path: Path) -> None:
    source = tmp_path / "episode_chus字幕.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\t中文\n"
        "Separate English line #\n\n"
        "00:00:02:00\t00:00:04:00\t中文//Inline English line #\n",
        encoding="utf-8",
    )
    output = tmp_path / "subtitle.docx"

    generate_docs(source, output_path=output)

    assert [paragraph.text for paragraph in Document(output).paragraphs] == [
        "00:00:00:00\t00:00:02:00\t中文",
        "Separate English line",
        "00:00:02:00\t00:00:04:00\t中文",
        "Inline English line",
    ]


def test_subtitle_highlights_blocks_without_inline_translation(tmp_path: Path) -> None:
    source = tmp_path / "episode_chus字幕.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\t中文\n"
        "Separate English\n\n"
        "00:00:02:00\t00:00:04:00\t中文//Inline English\n",
        encoding="utf-8",
    )
    output = tmp_path / "subtitle.docx"

    generate_docs(source, output_path=output)

    paragraphs = Document(output).paragraphs
    assert all(
        run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        for paragraph in paragraphs[:2]
        for run in paragraph.runs
    )
    assert all(
        run.font.highlight_color is None
        for paragraph in paragraphs[2:]
        for run in paragraph.runs
    )


def test_super_highlights_blocks_without_inline_translation(tmp_path: Path) -> None:
    source = tmp_path / "episode_super.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\n"
        "中文\n"
        "//Separate English\n\n"
        "00:00:02:00\t00:00:04:00\t//Inline English\n",
        encoding="utf-8",
    )
    output = tmp_path / "super.docx"

    generate_docs(source, output_path=output)

    paragraphs = Document(output).paragraphs
    assert all(
        run.font.highlight_color == WD_COLOR_INDEX.YELLOW
        for paragraph in paragraphs[:3]
        for run in paragraph.runs
    )
    assert paragraphs[3].text == ""
    assert all(
        run.font.highlight_color is None
        for paragraph in paragraphs[4:]
        for run in paragraph.runs
    )


def test_super_generation_preserves_notes_and_existing_english(tmp_path: Path) -> None:
    source = tmp_path / "episode_super.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\n"
        "風月同天\n"
        "//Sharing the Same Sky.\n\n"
        "(以下不用翻譯)\n"
        "歷史活動\n",
        encoding="utf-8",
    )
    output = tmp_path / "super.docx"

    with pytest.warns(UserWarning, match="periods or parentheses"):
        generate_docs(source, output_path=output)

    paragraphs = [paragraph.text for paragraph in Document(output).paragraphs]
    assert paragraphs == [
        "00:00:00:00\t00:00:02:00",
        "風月同天",
        "Sharing the Same Sky.",
        "",
        "(以下不用翻譯)",
        "歷史活動",
    ]


def test_super_generation_uses_official_program_title(tmp_path: Path) -> None:
    source = tmp_path / "episode_super.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\n風月同天\n",
        encoding="utf-8",
    )
    output = tmp_path / "super.docx"

    generate_docs(source, output_path=output)

    assert [paragraph.text for paragraph in Document(output).paragraphs] == [
        "00:00:00:00\t00:00:02:00",
        "風月同天",
        "Sharing the Same Sky",
    ]


def test_explicit_super_as_subtitle_note_uses_cyan(tmp_path: Path) -> None:
    source = tmp_path / "episode_super.txt"
    source.write_text(
        "(本段super以字幕方式呈現，已copy到字幕檔)\n"
        "00:00:00:00\t00:00:02:00\n"
        "字幕式SUPER\n",
        encoding="utf-8",
    )
    output = tmp_path / "super.docx"

    generate_docs(source, output_path=output)

    doc = Document(output)
    for paragraph in doc.paragraphs:
        assert paragraph.runs
        assert all(
            run.font.highlight_color == WD_COLOR_INDEX.TURQUOISE
            for run in paragraph.runs
        )


def test_shared_highlight_helper_preserves_cyan_and_adds_yellow(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "episode_super.txt"
    source.write_text(
        "(本段super以字幕方式呈現，已copy到字幕檔)\n"
        "00:00:00:00\t00:00:01:00\n"
        "明確標記\n\n"
        "00:00:02:00\t00:00:03:00\n"
        "一般內容\n",
        encoding="utf-8",
    )
    output = tmp_path / "super.docx"
    highlighted: list[tuple[str, WD_COLOR_INDEX]] = []
    shared_helper = generate_docs_module.apply_highlight_to_runs

    def record_highlight(paragraph, *, highlight_color) -> None:
        highlighted.append((paragraph.text, highlight_color))
        shared_helper(paragraph, highlight_color=highlight_color)

    monkeypatch.setattr(
        generate_docs_module,
        "apply_highlight_to_runs",
        record_highlight,
    )

    generate_docs(source, output_path=output)

    assert highlighted == [
        ("(本段super以字幕方式呈現，已copy到字幕檔)", WD_COLOR_INDEX.TURQUOISE),
        ("00:00:00:00\t00:00:01:00", WD_COLOR_INDEX.TURQUOISE),
        ("明確標記", WD_COLOR_INDEX.TURQUOISE),
        ("", WD_COLOR_INDEX.TURQUOISE),
        ("00:00:02:00\t00:00:03:00", WD_COLOR_INDEX.YELLOW),
        ("一般內容", WD_COLOR_INDEX.YELLOW),
    ]


def test_generated_document_is_valid_ooxml(tmp_path: Path) -> None:
    source = tmp_path / "episode_chus字幕.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\t中文//English\n",
        encoding="utf-8",
    )
    output = tmp_path / "subtitle.docx"

    generate_docs(source, output_path=output)

    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
        assert "word/document.xml" in package.namelist()


def test_generation_uses_synchronized_project_template_styles(tmp_path: Path) -> None:
    source = tmp_path / "episode_chus字幕.txt"
    source.write_text(
        "00:00:00:00\t00:00:02:00\t中文//English\n",
        encoding="utf-8",
    )
    output = tmp_path / "subtitle.docx"

    generate_docs(source, output_path=output)

    doc = Document(output)
    template = Document("templates/news_template.docx")
    assert doc.sections[0].page_width == template.sections[0].page_width
    assert doc.sections[0].page_height == template.sections[0].page_height
    assert doc.paragraphs[0].paragraph_format.line_spacing is None
    run = doc.paragraphs[0].runs[0]
    fonts = run._element.find("w:rPr/w:rFonts", run._element.nsmap)
    assert run.font.size.pt == BODY_TEXT_SIZE_PT
    assert fonts.get(qn("w:ascii")) == DEFAULT_DOCX_ASCII_FONT_NAME
    assert fonts.get(qn("w:eastAsia")) == DEFAULT_DOCX_EAST_ASIA_FONT_NAME
